"""Tests for the mapping reporter."""

from __future__ import annotations

import csv
import io
import json
import os
from dataclasses import dataclass

import pytest

from audit_logger import AuditLogger
from mapping_reporter import MappingReporter
from models import (
    AccountAssignmentRecord,
    ApplicationResult,
    EntitlementCreationResult,
    Inventory,
    PermissionSetRecord,
    RoleCreationResult,
)


@dataclass
class _Cfg:
    audit_to_cloudwatch: bool = False
    cloudwatch_log_group: str | None = None
    audit_to_file: bool = False
    audit_file_path: str | None = None
    workers: int = 1
    region: str = "us-east-1"
    run_id: str = "abc123"
    mapping_output_format: str = "JSON"
    mapping_output_path: str | None = None


def _audit():
    return AuditLogger("abc123", "caller", _Cfg(), stdout=io.StringIO(), stderr=io.StringIO())


def _inv(assignments):
    return Inventory(
        hub_account_id="111111111111",
        idc_instance_arn="arn:aws:sso:::instance/ssoins-x",
        identity_store_id="d-abc",
        permission_sets=(
            PermissionSetRecord(
                arn="arn:aws:sso:::permissionSet/ssoins-x/ps-1",
                name="Admin",
                description="",
                session_duration="PT1H",
                inline_policy=None,
                aws_managed_policy_arns=(),
                customer_managed_policy_references=(),
                permission_boundary=None,
            ),
        ),
        assignments=tuple(assignments),
        run_id="abc123",
        captured_at="2025-01-01T00:00:00+00:00",
    )


# ── Property 16: row count equals assignments ────────────────────────────────
def test_property_16_row_count_equals_assignments(tmp_path):
    """Property 16: Mapping report row count equals assignments.

    Validates: Requirement 14.2
    """
    assignments = [
        AccountAssignmentRecord(
            permission_set_arn="arn:aws:sso:::permissionSet/ssoins-x/ps-1",
            account_id="222222222222",
            principal_type="USER",
            principal_id=f"user-{i}",
            principal_display_name=f"User {i}",
        )
        for i in range(7)
    ]
    inventory = _inv(assignments)
    role_results = [
        RoleCreationResult(
            permission_set_arn="arn:aws:sso:::permissionSet/ssoins-x/ps-1",
            account_id="222222222222",
            role_name="AAM-Admin",
            role_arn="arn:aws:iam::222222222222:role/aam/AAM-Admin",
            status="CREATED",
        )
    ]
    application = ApplicationResult(
        application_arn="arn:aws:account-access:us-east-1:111111111111:application/app-1",
        status="SUPPLIED",
        idc_instance_arn=inventory.idc_instance_arn,
    )
    entitlement_results = [
        EntitlementCreationResult(
            application_arn=application.application_arn,
            permission_set_arn="arn:aws:sso:::permissionSet/ssoins-x/ps-1",
            account_id="222222222222",
            principal_type="USER",
            principal_id=f"user-{i}",
            role_arn="arn:aws:iam::222222222222:role/aam/AAM-Admin",
            entitlement_id=f"ent-{i}",
            status="CREATED",
        )
        for i in range(7)
    ]
    cfg = _Cfg(mapping_output_path=str(tmp_path / "out.json"))
    reporter = MappingReporter(cfg, _audit())
    rows = reporter.build_rows(inventory, role_results, application, entitlement_results)
    assert len(rows) == 7


# ── Property 17: row status accuracy ─────────────────────────────────────────
def test_property_17_row_status_accuracy(tmp_path):
    """Property 17: Mapping row status accuracy.

    Validates: Requirement 14.3
    """
    a = AccountAssignmentRecord(
        permission_set_arn="arn:aws:sso:::permissionSet/ssoins-x/ps-1",
        account_id="222222222222",
        principal_type="GROUP",
        principal_id="g-1",
        principal_display_name="Engineering",
    )
    inventory = _inv([a])
    role = RoleCreationResult(
        permission_set_arn=a.permission_set_arn,
        account_id=a.account_id,
        role_name="AAM-Admin",
        role_arn="arn:aws:iam::222222222222:role/aam/AAM-Admin",
        status="CREATED",
    )
    app_ok = ApplicationResult(
        application_arn="arn:aws:account-access:us-east-1:111111111111:application/app-1",
        status="SUPPLIED",
        idc_instance_arn=inventory.idc_instance_arn,
    )
    ent_existing = EntitlementCreationResult(
        application_arn=app_ok.application_arn,
        permission_set_arn=a.permission_set_arn,
        account_id=a.account_id,
        principal_type="GROUP",
        principal_id="g-1",
        role_arn=role.role_arn,
        entitlement_id="ent-existing",
        status="EXISTING",
    )
    rows = MappingReporter(_Cfg(), _audit()).build_rows(inventory, [role], app_ok, [ent_existing])
    assert rows[0].status == "EXISTING"

    # When the application failed, every row becomes FAILED
    app_failed = ApplicationResult(
        application_arn=None,
        status="FAILED",
        idc_instance_arn=inventory.idc_instance_arn,
        error_detail="boom",
    )
    rows = MappingReporter(_Cfg(), _audit()).build_rows(inventory, [role], app_failed, [ent_existing])
    assert rows[0].status == "FAILED"


# ── Property 18: empty mapping skips file write ───────────────────────────────
def test_property_18_empty_mapping_skips_file(tmp_path):
    """Property 18: Empty mapping skips file write.

    Validates: Requirement 14.10
    """
    inventory = _inv([])
    application = ApplicationResult(
        application_arn="app-arn",
        status="SUPPLIED",
        idc_instance_arn=inventory.idc_instance_arn,
    )
    cfg = _Cfg(mapping_output_path=str(tmp_path / "out.json"))
    reporter = MappingReporter(cfg, _audit())
    rows = reporter.build_rows(inventory, [], application, [])
    out = reporter.write(rows)
    assert out is None
    assert not os.path.exists(cfg.mapping_output_path)


# ── Format dispatch unit tests ────────────────────────────────────────────────
def _seed_for_format(tmp_path, fmt: str):
    a = AccountAssignmentRecord(
        permission_set_arn="arn:aws:sso:::permissionSet/ssoins-x/ps-1",
        account_id="222222222222",
        principal_type="USER",
        principal_id="user-1",
        principal_display_name="Bob",
    )
    inventory = _inv([a])
    role = RoleCreationResult(
        permission_set_arn=a.permission_set_arn,
        account_id=a.account_id,
        role_name="AAM-Admin",
        role_arn="arn:aws:iam::222222222222:role/aam/AAM-Admin",
        status="CREATED",
    )
    application = ApplicationResult(
        application_arn="arn:aws:account-access:us-east-1:111111111111:application/app-1",
        status="SUPPLIED",
        idc_instance_arn=inventory.idc_instance_arn,
    )
    ent = EntitlementCreationResult(
        application_arn=application.application_arn,
        permission_set_arn=a.permission_set_arn,
        account_id=a.account_id,
        principal_type="USER",
        principal_id="user-1",
        role_arn=role.role_arn,
        entitlement_id="ent-1",
        status="CREATED",
    )
    ext = fmt.lower()
    out = str(tmp_path / f"map.{ext}")
    cfg = _Cfg(mapping_output_format=fmt, mapping_output_path=out)
    reporter = MappingReporter(cfg, _audit())
    rows = reporter.build_rows(inventory, [role], application, [ent])
    return reporter, rows, out


def test_csv_output(tmp_path):
    reporter, rows, out = _seed_for_format(tmp_path, "CSV")
    reporter.write(rows)
    with open(out, encoding="utf-8") as f:
        lines = f.read().splitlines()
    assert lines[0].startswith("principal_type,principal_display_name")
    assert "Bob" in lines[1]


def test_json_output(tmp_path):
    reporter, rows, out = _seed_for_format(tmp_path, "JSON")
    reporter.write(rows)
    with open(out, encoding="utf-8") as f:
        data = json.load(f)
    assert isinstance(data, list)
    assert data[0]["principal_display_name"] == "Bob"


def test_xlsx_output(tmp_path):
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    reporter, rows, out = _seed_for_format(tmp_path, "XLSX")
    reporter.write(rows)
    wb = load_workbook(out)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert headers[0] == "principal_type"
    body = [c.value for c in ws[2]]
    assert "Bob" in body


def test_default_filename_uses_run_id(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    a = AccountAssignmentRecord(
        permission_set_arn="arn:aws:sso:::permissionSet/ssoins-x/ps-1",
        account_id="222222222222",
        principal_type="USER",
        principal_id="user-1",
        principal_display_name="Bob",
    )
    inventory = _inv([a])
    role = RoleCreationResult(
        permission_set_arn=a.permission_set_arn,
        account_id=a.account_id,
        role_name="AAM-Admin",
        role_arn="arn:aws:iam::222222222222:role/aam/AAM-Admin",
        status="CREATED",
    )
    application = ApplicationResult(
        application_arn="app-1",
        status="SUPPLIED",
        idc_instance_arn=inventory.idc_instance_arn,
    )
    cfg = _Cfg(mapping_output_format="CSV", run_id="rid42")
    reporter = MappingReporter(cfg, _audit())
    rows = reporter.build_rows(inventory, [role], application, [])
    out = reporter.write(rows)
    assert out is not None
    assert os.path.basename(out) == "mapping_rid42.csv"
